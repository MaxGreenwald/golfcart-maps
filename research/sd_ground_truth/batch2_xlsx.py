#!/usr/bin/env python3
"""streetview-batch2.csv -> streetview-batch2.xlsx in the batch 1 sheet layout (Roads with agree/resolved formulas, How to check, Summary)."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
HERE = os.path.dirname(os.path.abspath(__file__))
R = list(csv.DictReader(open(os.path.join(HERE, "streetview-batch2.csv")))); COLS = list(R[0].keys()); N = len(R) + 1
wb = Workbook(); ws = wb.active; ws.title = "Roads"; ws.append(COLS)
ci = {c: i + 1 for i, c in enumerate(COLS)}; Lt = lambda c: ws.cell(1, ci[c]).column_letter
for i, r in enumerate(R, start=2):
    for c in COLS:
        v = r[c]; ws.cell(i, ci[c], int(v) if v.isdigit() and c in ("our_speed_mph", "routes_using", "length_m", "our_gis_survey_mph") else v)
    P, S, A = f"{Lt('posted_speed_seen_mph')}{i}", f"{Lt('our_gis_survey_mph')}{i}", f"{Lt('agree_YN')}{i}"
    ws[A] = f'=IF(OR({P}="",{S}=""),"",IF(VALUE({P})=VALUE({S}),"Y","N"))'; ws[f"{Lt('resolved_mph')}{i}"] = f'=IF({A}="Y",{S},"")'
VOL = ("posted_speed_seen_mph", "sign_location_note", "imagery_date", "confidence", "status", "checker_initials", "notes")
for c in range(1, len(COLS) + 1):
    ws.cell(1, c).font = Font(bold=True); ws.cell(1, c).fill = PatternFill("solid", fgColor="E6FBFC" if COLS[c-1] in VOL else "FFF4CC")
ws.freeze_panes = "D2"
dv = DataValidation(type="list", formula1='"pending,done,can\'t tell"'); ws.add_data_validation(dv); dv.add(f"{Lt('status')}2:{Lt('status')}{N}")
dc = DataValidation(type="list", formula1='"high,medium,low"'); ws.add_data_validation(dc); dc.add(f"{Lt('confidence')}2:{Lt('confidence')}{N}")
for c, w in {"road_name": 24, "from_cross_street": 20, "to_cross_street": 20, "street_view_link": 18, "map_link": 14, "why_in_batch": 50, "notes": 30}.items(): ws.column_dimensions[Lt(c)].width = w
h = wb.create_sheet("How to check")
for line in ["Street View batch 2: 1,000 stretches, San Diego first. Rows are in priority order within each region: RED ISLAND rows first (a short piece of a road shows as not allowed while the same road is 35 mph or less on both sides; tell us the posted limit on that piece), then roads whose limit we only guessed or took from one unverified source, right at the 25 mph (golf cart) or 35 mph (LSV) line.",
             "Same as batch 1: fill only the blue columns (posted_speed_seen_mph … notes). If no speed sign is visible on the stretch, set status to \"can't tell\" and move on (5 minute cap per row). Row order is priority order, so stopping partway is fine.", ""] + open("/workspace/gcm/gtm/streetview-checker-guide.md").read().replace("Golf Cart Map ", "Golf Cart Maps ").replace("Golf Cart Map\n", "Golf Cart Maps\n").splitlines():
    h.append([line])
h.column_dimensions["A"].width = 160
s = wb.create_sheet("Summary"); P, S, A, ST = Lt("posted_speed_seen_mph"), Lt("our_gis_survey_mph"), Lt("agree_YN"), Lt("status")
for k, f in [("metric", "value"), ("total roads", f"=COUNTA(Roads!A2:A{N})"), ("done", f'=COUNTIF(Roads!{ST}2:{ST}{N},"done")'),
             ("can't tell", f"=COUNTIF(Roads!{ST}2:{ST}{N},\"can't tell\")"), ("pending", f'=COUNTIF(Roads!{ST}2:{ST}{N},"pending")'),
             ("percent done", "=IF(B2=0,0,(B3+B4)/B2)"), ("compared (sign read and we have a GIS/survey value)", f'=COUNTIF(Roads!{A}2:{A}{N},"Y")+COUNTIF(Roads!{A}2:{A}{N},"N")'),
             ("agreed (Y)", f'=COUNTIF(Roads!{A}2:{A}{N},"Y")'), ("disagreed (N)", f'=COUNTIF(Roads!{A}2:{A}{N},"N")'), ("% agreed", '=IF(B7=0,"",B8/B7)'),
             ("mismatched vs our map (done rows)", f'=SUMPRODUCT((Roads!{ST}2:{ST}{N}="done")*(Roads!{P}2:{P}{N}<>"")*(Roads!{P}2:{P}{N}<>Roads!G2:G{N}))'),
             ("red island rows", f'=COUNTIF(Roads!{Lt("why_in_batch")}2:{Lt("why_in_batch")}{N},"RED ISLAND*")')]:
    s.append([k, f])
s["B6"].number_format = "0.0%"; s["B10"].number_format = "0.0%"; s.column_dimensions["A"].width = 55
wb.save(os.path.join(HERE, "streetview-batch2.xlsx")); print("ok", len(R))
